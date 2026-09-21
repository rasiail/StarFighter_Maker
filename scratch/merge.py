import openpyxl

wb = openpyxl.load_workbook('balance/cards.xlsx')
cards_sheet = wb['Cards']
effects_sheet = wb['Effects']

# Read all effects
effects = {}
# effects_sheet row 4 is headers, row 5 is ko descriptions, row 6+ is data
for row in effects_sheet.iter_rows(min_row=6, values_only=True):
    if not row[0]: continue
    cid = row[0]
    if cid not in effects: effects[cid] = []
    effects[cid].append(row)

# Add headers to Cards
cards_sheet.cell(row=4, column=8).value = 'e1_key'
cards_sheet.cell(row=4, column=9).value = 'e1_op'
cards_sheet.cell(row=4, column=10).value = 'e1_val'
cards_sheet.cell(row=4, column=11).value = 'e1_unit'
cards_sheet.cell(row=4, column=12).value = 'e2_key'
cards_sheet.cell(row=4, column=13).value = 'e2_op'
cards_sheet.cell(row=4, column=14).value = 'e2_val'
cards_sheet.cell(row=4, column=15).value = 'e2_unit'
cards_sheet.cell(row=4, column=16).value = 'e3_key'
cards_sheet.cell(row=4, column=17).value = 'e3_op'
cards_sheet.cell(row=4, column=18).value = 'e3_val'
cards_sheet.cell(row=4, column=19).value = 'e3_unit'

# Add Korean descriptions to row 5
cards_sheet.cell(row=5, column=8).value = '효과1 키'
cards_sheet.cell(row=5, column=9).value = '효과1 연산'
cards_sheet.cell(row=5, column=10).value = '효과1 수치'
cards_sheet.cell(row=5, column=11).value = '효과1 단위'
cards_sheet.cell(row=5, column=12).value = '효과2 키'
cards_sheet.cell(row=5, column=13).value = '효과2 연산'
cards_sheet.cell(row=5, column=14).value = '효과2 수치'
cards_sheet.cell(row=5, column=15).value = '효과2 단위'
cards_sheet.cell(row=5, column=16).value = '효과3 키'
cards_sheet.cell(row=5, column=17).value = '효과3 연산'
cards_sheet.cell(row=5, column=18).value = '효과3 수치'
cards_sheet.cell(row=5, column=19).value = '효과3 단위'

formulas = {
    'mobility': '="피치·롤·요 최대 회전속도 +" & (J{r}*100) & "%"',
    'stability': '="회전 정리·역입력 응답 +" & (J{r}*100) & "%"',
    'speed': '="순항 +" & J{r} & " / 최고 +" & N{r} & " kts · 가속 +" & R{r}',
    'defense': '="최대 체력 +" & J{r} & " · 증가분 회복"',
    'power': '="기본 무기 피해 배율 +" & (J{r}*100) & "%"',
    'control': '="락온 거리 +" & (J{r}*100) & "% · 유도 선회 +" & (N{r}*100) & "%"',
    'standardRack': '="표준 미사일 탄창 +" & J{r} & "발"',
    'multiRack': '="멀티 미사일 탄창 +" & J{r} & "발"',
    'reload': '="두 미사일의 기본 재장전 시간 -" & (J{r}*100) & "%"',
    'warhead': '="화력 보정 후 무기 피해 +" & (J{r}*100) & "% · 화력 3 필요"',
    'guidance': '="미사일 유도 선회 성능 +" & (J{r}*100) & "% · 관제력 3 필요"',
    'repair': '="체력 " & (J{r}*100) & "% 회복 · 점수 +" & N{r}',
    'multiSalvo': '="멀티 동시 락온·발사 +" & J{r} & " (4→" & (4+J{r}) & "→" & (4+J{r}*2) & "발) · 관제력 2 필요"'
}

for row_idx in range(6, cards_sheet.max_row + 1):
    cid = cards_sheet.cell(row=row_idx, column=1).value
    if not cid: continue
    
    # Write formulas
    if cid in formulas:
        cards_sheet.cell(row=row_idx, column=7).value = formulas[cid].format(r=row_idx)

    # Write effects
    c_effects = effects.get(cid, [])
    for i, eff in enumerate(c_effects):
        if i >= 3: break
        col_offset = 8 + i * 4
        cards_sheet.cell(row=row_idx, column=col_offset).value = eff[1]
        cards_sheet.cell(row=row_idx, column=col_offset+1).value = eff[2]
        cards_sheet.cell(row=row_idx, column=col_offset+3).value = eff[4]
        # value is at eff[3]. If it's 0.08, let's just write the raw float.
        cards_sheet.cell(row=row_idx, column=col_offset+2).value = eff[3]

wb.remove(effects_sheet)
wb.save('balance/cards.xlsx')
print('Merge complete')

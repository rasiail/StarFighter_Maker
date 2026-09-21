import openpyxl

wb = openpyxl.load_workbook('balance/levels.xlsx', data_only=False)
for name in wb.sheetnames:
    print("Sheet:", name)
    ws = wb[name]
    for r in range(1, 15):
        row_vals = [ws.cell(r, c).value for c in range(1, 10)]
        if any(row_vals):
            print(f"Row {r}: {row_vals}")

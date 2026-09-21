import openpyxl

wb = openpyxl.load_workbook('balance/levels.xlsx')
ws = wb['Levels']

print("Header row 4:", [ws.cell(4, c).value for c in range(1, 5)])
print("Header row 5:", [ws.cell(5, c).value for c in range(1, 5)])
for r in range(6, 15):
    print(f"Row {r} (Lv {ws.cell(r,1).value}): xp_to_next={ws.cell(r,2).value}, cum={ws.cell(r,3).value}")

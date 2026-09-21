import os
import openpyxl
import win32com.client

file_path = os.path.abspath('balance/levels.xlsx')

# 1. Update values via openpyxl
wb = openpyxl.load_workbook(file_path)
ws = wb['Levels']

# Row 6 is Level 1, Row 7 is Level 2, ...
new_early_xp = {
    6: 60,   # Lv 1
    7: 150,  # Lv 2
    8: 300,  # Lv 3
    9: 500,  # Lv 4
}

for row, xp in new_early_xp.items():
    ws.cell(row, 2).value = xp

# Fix cumulative_xp_at_start formulas for all level rows (Row 6 to 65)
ws.cell(6, 3).value = 0
for r in range(7, 66):
    ws.cell(r, 3).value = f'=C{r-1}+B{r-1}'

wb.save(file_path)
print("openpyxl update complete.")

# 2. Open and save via Excel COM to evaluate and cache formula values
excel = win32com.client.DispatchEx('Excel.Application')
excel.Visible = False
excel.DisplayAlerts = False
try:
    workbook = excel.Workbooks.Open(file_path)
    workbook.Save()
    workbook.Close()
    print("Excel COM save complete.")
finally:
    excel.Quit()
